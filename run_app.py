#!/usr/bin/env python3
"""
Cyberpunk Test Automation Control Panel
Run: python run_app.py
"""
import os, io, re, json, base64, datetime, textwrap, threading, subprocess, webbrowser
from http.server import HTTPServer, BaseHTTPRequestHandler
from urllib.parse import parse_qs, urlparse

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

ROBOT_DIR   = os.path.dirname(os.path.abspath(__file__))
ADB_PATH    = os.path.join(ROBOT_DIR, "adb.exe")
PORT        = 8080

DEVICE_NAMES = {
    "RFCN80A7HYJ": "Galaxy Note 20",
}

run_state = {"running": False, "output": [], "returncode": None, "done": False}
state_lock = threading.Lock()

# ═══════════════════════════════════════════════════════════════════════════════
#  MAIN CONTROL PANEL HTML
# ═══════════════════════════════════════════════════════════════════════════════
HTML = r"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>TESTAUTOMATION // CONTROL PANEL</title>
<style>
  :root {
    --bg:#040408; --bg1:#080810; --bg2:#0d0d1a; --bg3:#111122;
    --cyan:#00f5ff; --cyan2:#00c8d4; --magenta:#ff00ff; --magenta2:#cc00cc;
    --green:#00ff41; --red:#ff2244; --amber:#ffaa00;
    --dim:#334; --border:#1a1a3a; --text:#c8c8e8; --text-dim:#556;
  }
  *,*::before,*::after{box-sizing:border-box;margin:0;padding:0}
  body{font-family:'Courier New','Consolas',monospace;background:var(--bg);color:var(--text);min-height:100vh;overflow-x:hidden}
  body::before{content:'';position:fixed;inset:0;background-image:linear-gradient(rgba(0,245,255,.03) 1px,transparent 1px),linear-gradient(90deg,rgba(0,245,255,.03) 1px,transparent 1px);background-size:40px 40px;pointer-events:none;z-index:0}
  body::after{content:'';position:fixed;inset:0;background:repeating-linear-gradient(0deg,transparent,transparent 2px,rgba(0,0,0,.15) 2px,rgba(0,0,0,.15) 4px);pointer-events:none;z-index:0}
  .wrap{position:relative;z-index:1;max-width:1100px;margin:0 auto;padding:24px 20px 40px}
  header{display:flex;align-items:center;justify-content:space-between;border-bottom:1px solid var(--cyan);padding-bottom:14px;margin-bottom:28px;gap:16px;flex-wrap:wrap}
  .logo{font-size:1.25rem;font-weight:700;letter-spacing:.2em;color:var(--cyan);text-shadow:0 0 12px var(--cyan),0 0 30px rgba(0,245,255,.4)}
  .logo span{color:var(--magenta);text-shadow:0 0 12px var(--magenta)}
  .header-meta{display:flex;gap:10px;align-items:center;flex-wrap:wrap}
  .pill{font-size:.72rem;letter-spacing:.1em;padding:3px 10px;border:1px solid var(--border);border-radius:2px;color:var(--text-dim)}
  .pill.online{border-color:var(--green);color:var(--green);text-shadow:0 0 6px var(--green)}
  .btn-gen{font-family:inherit;font-size:.72rem;letter-spacing:.12em;padding:5px 14px;border:1px solid var(--magenta);border-radius:2px;color:var(--magenta);background:transparent;cursor:pointer;text-decoration:none;transition:all .15s;display:inline-block}
  .btn-gen:hover{background:rgba(255,0,255,.08);box-shadow:0 0 12px rgba(255,0,255,.3)}
  .grid{display:grid;grid-template-columns:280px 1fr;gap:18px;margin-bottom:18px}
  @media(max-width:700px){.grid{grid-template-columns:1fr}}
  .card{background:var(--bg2);border:1px solid var(--border);border-radius:4px;padding:20px;position:relative}
  .card::before{content:'';position:absolute;top:-1px;left:-1px;width:12px;height:12px;border-top:2px solid var(--cyan);border-left:2px solid var(--cyan)}
  .card::after{content:'';position:absolute;bottom:-1px;right:-1px;width:12px;height:12px;border-bottom:2px solid var(--cyan);border-right:2px solid var(--cyan)}
  .card-title{font-size:.68rem;letter-spacing:.2em;color:var(--cyan);text-shadow:0 0 8px var(--cyan);margin-bottom:16px;padding-bottom:8px;border-bottom:1px solid var(--border)}
  #device-list{font-size:.82rem;line-height:2}
  .device-row{display:flex;align-items:center;gap:8px}
  .d-dot{width:8px;height:8px;border-radius:50%;flex-shrink:0}
  .d-dot.online{background:var(--green);box-shadow:0 0 6px var(--green)}
  .d-dot.offline{background:var(--red)}
  .device-name{color:var(--cyan2)}
  .device-status{color:var(--text-dim);font-size:.72rem;margin-left:auto}
  .adb-ver{color:var(--text-dim);font-size:.72rem;margin-bottom:12px}
  label{font-size:.72rem;letter-spacing:.12em;color:var(--text-dim);display:block;margin-bottom:5px}
  select,input[type=text]{width:100%;background:var(--bg1);border:1px solid var(--border);color:var(--cyan);font-family:inherit;font-size:.82rem;padding:8px 10px;border-radius:2px;outline:none;margin-bottom:14px;transition:border-color .2s,box-shadow .2s}
  select:focus,input[type=text]:focus{border-color:var(--cyan);box-shadow:0 0 8px rgba(0,245,255,.25)}
  .btn-row{display:flex;gap:10px;margin-top:6px;flex-wrap:wrap}
  button{font-family:inherit;font-size:.78rem;letter-spacing:.12em;padding:9px 22px;border:1px solid;border-radius:2px;cursor:pointer;transition:all .15s;position:relative;overflow:hidden}
  button::after{content:'';position:absolute;inset:0;background:rgba(255,255,255,.04);opacity:0;transition:opacity .15s}
  button:hover::after{opacity:1}
  .btn-run{background:transparent;border-color:var(--cyan);color:var(--cyan);text-shadow:0 0 8px var(--cyan);box-shadow:0 0 12px rgba(0,245,255,.2),inset 0 0 12px rgba(0,245,255,.05)}
  .btn-run:hover:not(:disabled){background:rgba(0,245,255,.08);box-shadow:0 0 20px rgba(0,245,255,.4),inset 0 0 16px rgba(0,245,255,.1)}
  .btn-run:disabled{border-color:var(--dim);color:var(--dim);box-shadow:none;cursor:not-allowed}
  .btn-abort{background:transparent;border-color:var(--red);color:var(--red);display:none}
  .btn-abort:hover{background:rgba(255,34,68,.1);box-shadow:0 0 12px rgba(255,34,68,.3)}
  .btn-abort.visible{display:inline-block}
  .btn-ghost{background:transparent;border-color:var(--border);color:var(--text-dim);display:none}
  .btn-ghost:hover{border-color:var(--magenta);color:var(--magenta)}
  .btn-ghost.visible{display:inline-block}
  .btn-refresh{background:transparent;border-color:var(--border);color:var(--text-dim);font-size:.7rem;padding:6px 14px}
  .btn-refresh:hover{border-color:var(--cyan2);color:var(--cyan2)}
  .status-bar{display:flex;align-items:center;gap:10px;margin-bottom:10px;font-size:.78rem;letter-spacing:.08em}
  .s-dot{width:9px;height:9px;border-radius:50%;background:var(--dim);flex-shrink:0;transition:background .3s}
  .s-dot.running{background:var(--amber);box-shadow:0 0 8px var(--amber);animation:blink 1s infinite}
  .s-dot.pass{background:var(--green);box-shadow:0 0 8px var(--green)}
  .s-dot.fail{background:var(--red);box-shadow:0 0 8px var(--red)}
  @keyframes blink{0%,100%{opacity:1}50%{opacity:.25}}
  .term-card{background:var(--bg1);border:1px solid var(--border);border-radius:4px;overflow:hidden;position:relative}
  .term-card::before{content:'';position:absolute;top:-1px;left:-1px;width:12px;height:12px;border-top:2px solid var(--magenta);border-left:2px solid var(--magenta)}
  .term-card::after{content:'';position:absolute;bottom:-1px;right:-1px;width:12px;height:12px;border-bottom:2px solid var(--magenta);border-right:2px solid var(--magenta)}
  .term-header{background:var(--bg2);border-bottom:1px solid var(--border);padding:8px 16px;display:flex;align-items:center;gap:12px}
  .term-title{font-size:.68rem;letter-spacing:.2em;color:var(--magenta);text-shadow:0 0 8px var(--magenta)}
  .term-dots{display:flex;gap:6px;margin-left:auto}
  .term-dot{width:8px;height:8px;border-radius:50%}
  #output{min-height:340px;max-height:480px;overflow-y:auto;padding:16px;font-size:.8rem;line-height:1.7;white-space:pre-wrap;word-break:break-all;scrollbar-width:thin;scrollbar-color:var(--border) transparent}
  #output::-webkit-scrollbar{width:6px}
  #output::-webkit-scrollbar-thumb{background:var(--border);border-radius:3px}
  .l-pass{color:var(--green);text-shadow:0 0 4px rgba(0,255,65,.5)}
  .l-fail{color:var(--red);text-shadow:0 0 4px rgba(255,34,68,.5)}
  .l-warn{color:var(--amber);text-shadow:0 0 4px rgba(255,170,0,.4)}
  .l-info{color:var(--cyan);text-shadow:0 0 4px rgba(0,245,255,.4)}
  .l-plain{color:#889}
  .l-sep{color:#334}
  .l-prompt::before{content:'> ';color:var(--magenta)}
  #placeholder{color:var(--text-dim);font-style:italic}
  .cursor{display:inline-block;width:8px;height:14px;background:var(--cyan);animation:blink 1s infinite;vertical-align:text-bottom;margin-left:2px}
  .footer{display:flex;align-items:center;gap:12px;margin-top:14px;font-size:.72rem;color:var(--text-dim);letter-spacing:.08em;flex-wrap:wrap}
  .footer-right{margin-left:auto;display:flex;gap:10px;align-items:center}
</style>
</head>
<body>
<div class="wrap">
  <header>
    <div class="logo">TEST<span>//</span>AUTOMATION</div>
    <div class="header-meta">
      <div class="pill">v1.0</div>
      <div class="pill">ROBOT FRAMEWORK</div>
      <div class="pill online">&#11044; ONLINE</div>
      <div class="pill" id="date-de" style="color:var(--amber);border-color:var(--amber);text-shadow:0 0 6px rgba(255,170,0,.5)"></div>
      <a href="/generator" class="btn-gen">&#9889; SCRIPT GENERATOR</a>
    </div>
  </header>

  <div class="grid">
    <div class="card">
      <div class="card-title">// DEVICE STATUS</div>
      <div class="adb-ver" id="adb-ver">ADB &middot;&middot;&middot;</div>
      <div id="device-list"><span style="color:var(--text-dim)">Scanning&middot;&middot;&middot;</span></div>
      <br>
      <button class="btn-refresh" onclick="refreshDevices()">&#8635;  REFRESH DEVICES</button>
    </div>

    <div class="card">
      <div class="card-title">// MISSION CONTROL</div>
      <label>TEST SUITE</label>
      <select id="suite"><option value="">Loading&hellip;</option></select>
      <label>PHONE NUMBER (for call tests)</label>
      <input type="text" id="phone" value="555" placeholder="+15551234567" />
      <div class="status-bar">
        <div class="s-dot" id="dot"></div>
        <span id="statusText">READY</span>
      </div>
      <div class="btn-row">
        <button class="btn-run"   id="runBtn"    onclick="runTest()">&#9654;  EXECUTE</button>
        <button class="btn-abort" id="abortBtn"  onclick="abortTest()">&#9632;  ABORT</button>
        <button class="btn-ghost" id="reportBtn" onclick="openReport()">&#128196;  REPORT</button>
        <button class="btn-ghost" id="logBtn"    onclick="openLog()">&#128196;  LOG</button>
      </div>
    </div>
  </div>

  <div class="term-card">
    <div class="term-header">
      <div class="term-title">// TERMINAL OUTPUT</div>
      <div class="term-dots">
        <div class="term-dot" style="background:#ff2244"></div>
        <div class="term-dot" style="background:#ffaa00"></div>
        <div class="term-dot" style="background:#00ff41"></div>
      </div>
    </div>
    <div id="output">
      <span id="placeholder">Awaiting execution command&middot;&middot;&middot;<span class="cursor"></span></span>
    </div>
  </div>

  <div class="footer">
    <span id="footerStatus">SYSTEM IDLE</span>
    <span>|</span>
    <span>PORT <span style="color:var(--cyan)">8080</span></span>
    <div class="footer-right" id="footer-time"></div>
  </div>
</div>

<script>
  function tick() {
    const d = new Date();
    document.getElementById('footer-time').textContent =
      d.toLocaleTimeString('en-GB', {hour12:false}) + '  UTC' +
      (d.getTimezoneOffset()<=0?'+':'-') +
      String(Math.abs(d.getTimezoneOffset()/60)).padStart(2,'0');
    document.getElementById('date-de').textContent =
      d.toLocaleDateString('de-DE', {day:'2-digit', month:'long', year:'numeric'});
  }
  setInterval(tick, 1000); tick();

  function esc(s){return String(s).replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;')}

  function colorLine(line) {
    if (/PASS|passed/i.test(line))  return '<span class="l-pass">'+esc(line)+'</span>';
    if (/FAIL|ERROR|error/i.test(line)) return '<span class="l-fail">'+esc(line)+'</span>';
    if (/WARN|warn/i.test(line))    return '<span class="l-warn">'+esc(line)+'</span>';
    if (/^={3,}/.test(line))        return '<span class="l-sep">'+esc(line)+'</span>';
    if (/Robot Framework|Running|Output|Report|Log|==/.test(line)) return '<span class="l-info">'+esc(line)+'</span>';
    return '<span class="l-plain l-prompt">'+esc(line)+'</span>';
  }

  async function loadScripts() {
    try {
      const res  = await fetch('/list-scripts');
      const data = await res.json();
      const sel  = document.getElementById('suite');
      sel.innerHTML = data.scripts.map(s => '<option value="'+esc(s)+'">'+esc(s)+'</option>').join('');
      // Pre-select from URL ?suite=
      const param = new URLSearchParams(window.location.search).get('suite');
      if (param) {
        const opt = Array.from(sel.options).find(o => o.value === param);
        if (opt) opt.selected = true;
      }
    } catch(e) {
      document.getElementById('suite').innerHTML = '<option value="call_test.robot">call_test.robot</option>';
    }
  }
  loadScripts();

  async function refreshDevices() {
    document.getElementById('device-list').innerHTML = '<span style="color:var(--text-dim)">Scanning&middot;&middot;&middot;</span>';
    document.getElementById('adb-ver').textContent = 'ADB &middot;&middot;&middot;';
    const res  = await fetch('/devices');
    const data = await res.json();
    document.getElementById('adb-ver').textContent = data.version || 'ADB not found';
    const list = document.getElementById('device-list');
    if (!data.devices || !data.devices.length) {
      list.innerHTML = '<span style="color:var(--text-dim)">No devices detected</span>'; return;
    }
    list.innerHTML = data.devices.map(d =>
      '<div class="device-row"><div class="d-dot '+(d.status==='device'?'online':'offline')+'"></div>'+
      '<span class="device-name">'+(d.name?esc(d.name):esc(d.serial))+'</span>'+
      '<span class="device-status" title="'+esc(d.serial)+'">'+esc(d.serial)+'  '+esc(d.status)+'</span></div>'
    ).join('');
  }
  refreshDevices();

  let polling = null, lineIndex = 0;

  async function runTest() {
    const suite = document.getElementById('suite').value;
    const phone = document.getElementById('phone').value.trim() || '555';
    document.getElementById('runBtn').disabled = true;
    document.getElementById('abortBtn').classList.add('visible');
    document.getElementById('reportBtn').classList.remove('visible');
    document.getElementById('logBtn').classList.remove('visible');
    document.getElementById('output').innerHTML = '';
    document.getElementById('dot').className = 's-dot running';
    document.getElementById('statusText').textContent = 'RUNNING  '+suite;
    document.getElementById('footerStatus').textContent = 'EXECUTING '+suite.toUpperCase();
    lineIndex = 0;
    await fetch('/run', {method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({suite,phone})});
    polling = setInterval(pollStatus, 600);
  }

  async function pollStatus() {
    const res  = await fetch('/status?from='+lineIndex);
    const data = await res.json();
    if (data.lines && data.lines.length) {
      const out = document.getElementById('output');
      data.lines.forEach(l => { out.innerHTML += colorLine(l)+'\n'; lineIndex++; });
      out.scrollTop = out.scrollHeight;
    }
    if (data.done) {
      clearInterval(polling);
      document.getElementById('runBtn').disabled = false;
      document.getElementById('abortBtn').classList.remove('visible');
      document.getElementById('reportBtn').classList.add('visible');
      document.getElementById('logBtn').classList.add('visible');
      if (data.returncode === 0) {
        document.getElementById('dot').className = 's-dot pass';
        document.getElementById('statusText').textContent = '&#10003; ALL TESTS PASSED';
        document.getElementById('footerStatus').textContent = 'LAST RUN: PASS';
      } else {
        document.getElementById('dot').className = 's-dot fail';
        document.getElementById('statusText').textContent = '&#10007; TESTS FAILED  (code '+data.returncode+')';
        document.getElementById('footerStatus').textContent = 'LAST RUN: FAIL';
      }
    }
  }

  async function abortTest() {
    await fetch('/abort',{method:'POST'});
    document.getElementById('statusText').textContent = 'ABORT SIGNAL SENT';
  }
  function openReport(){window.open('/report','_blank')}
  function openLog()   {window.open('/log','_blank')}
</script>
</body>
</html>
"""

# ═══════════════════════════════════════════════════════════════════════════════
#  SCRIPT GENERATOR PAGE HTML
# ═══════════════════════════════════════════════════════════════════════════════
GENERATOR_HTML = r"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>SCRIPT GENERATOR // TEST AUTOMATION</title>
<style>
  :root{
    --bg:#040408;--bg1:#080810;--bg2:#0d0d1a;--bg3:#111122;
    --cyan:#00f5ff;--cyan2:#00c8d4;--magenta:#ff00ff;--magenta2:#cc00cc;
    --green:#00ff41;--red:#ff2244;--amber:#ffaa00;--teal:#00c8d4;
    --dim:#334;--border:#1a1a3a;--text:#c8c8e8;--text-dim:#556;
  }
  *,*::before,*::after{box-sizing:border-box;margin:0;padding:0}
  body{font-family:'Courier New','Consolas',monospace;background:var(--bg);color:var(--text);min-height:100vh;overflow-x:hidden}
  body::before{content:'';position:fixed;inset:0;background-image:linear-gradient(rgba(0,245,255,.03) 1px,transparent 1px),linear-gradient(90deg,rgba(0,245,255,.03) 1px,transparent 1px);background-size:40px 40px;pointer-events:none;z-index:0}
  body::after{content:'';position:fixed;inset:0;background:repeating-linear-gradient(0deg,transparent,transparent 2px,rgba(0,0,0,.15) 2px,rgba(0,0,0,.15) 4px);pointer-events:none;z-index:0}
  .wrap{position:relative;z-index:1;max-width:1100px;margin:0 auto;padding:24px 20px 60px}

  /* Header */
  header{display:flex;align-items:center;justify-content:space-between;border-bottom:1px solid var(--cyan);padding-bottom:14px;margin-bottom:24px;gap:16px;flex-wrap:wrap}
  .logo{font-size:1.25rem;font-weight:700;letter-spacing:.2em;color:var(--cyan);text-shadow:0 0 12px var(--cyan),0 0 30px rgba(0,245,255,.4)}
  .logo span{color:var(--magenta);text-shadow:0 0 12px var(--magenta)}
  .back-link{color:var(--text-dim);text-decoration:none;font-size:.72rem;letter-spacing:.12em;border:1px solid var(--border);padding:4px 12px;border-radius:2px;transition:all .2s}
  .back-link:hover{color:var(--cyan);border-color:var(--cyan)}
  .header-meta{display:flex;gap:10px;align-items:center;flex-wrap:wrap}
  .pill{font-size:.72rem;letter-spacing:.1em;padding:3px 10px;border:1px solid var(--border);border-radius:2px;color:var(--text-dim)}
  .pill.online{border-color:var(--green);color:var(--green);text-shadow:0 0 6px var(--green)}

  /* Cards */
  .card{background:var(--bg2);border:1px solid var(--border);border-radius:4px;padding:22px;position:relative;margin-bottom:16px}
  .card::before{content:'';position:absolute;top:-1px;left:-1px;width:12px;height:12px;border-top:2px solid var(--cyan);border-left:2px solid var(--cyan)}
  .card::after{content:'';position:absolute;bottom:-1px;right:-1px;width:12px;height:12px;border-bottom:2px solid var(--cyan);border-right:2px solid var(--cyan)}
  .card-title{font-size:.68rem;letter-spacing:.2em;color:var(--cyan);text-shadow:0 0 8px var(--cyan);margin-bottom:16px;padding-bottom:8px;border-bottom:1px solid var(--border)}

  /* Upload zone */
  .drop-zone{border:2px dashed var(--border);border-radius:4px;padding:2.5rem 2rem;text-align:center;cursor:pointer;transition:all .2s;user-select:none}
  .drop-zone:hover,.drop-zone.over{border-color:var(--cyan);background:rgba(0,245,255,.03);box-shadow:0 0 20px rgba(0,245,255,.08) inset}
  .drop-icon{font-size:2rem;margin-bottom:.75rem;color:var(--cyan);opacity:.55;line-height:1}
  .drop-text{font-size:.85rem;letter-spacing:.15em;color:var(--text-dim);margin-bottom:.4rem}
  .drop-sub{font-size:.7rem;color:var(--text-dim);opacity:.6}

  /* Type selector */
  .type-row{display:flex;gap:20px;margin:14px 0;flex-wrap:wrap}
  .type-opt{display:flex;align-items:center;gap:8px;cursor:pointer;font-size:.78rem;letter-spacing:.08em;color:var(--text-dim);transition:color .2s}
  .type-opt input{accent-color:var(--cyan);cursor:pointer;width:14px;height:14px}
  .type-opt:hover{color:var(--cyan)}
  .type-opt input:checked+span{color:var(--cyan);text-shadow:0 0 6px var(--cyan)}

  /* File info */
  .file-info{display:flex;align-items:center;gap:10px;margin-top:12px;font-size:.8rem;padding:8px 12px;background:rgba(0,255,65,.05);border:1px solid rgba(0,255,65,.2);border-radius:2px}
  .file-name{color:var(--green);font-family:inherit}
  .file-size{color:var(--text-dim)}

  /* Buttons */
  .btn-row{display:flex;gap:10px;margin-top:14px;flex-wrap:wrap}
  button{font-family:inherit;font-size:.78rem;letter-spacing:.12em;padding:9px 20px;border:1px solid;border-radius:2px;cursor:pointer;transition:all .15s;position:relative;overflow:hidden}
  .btn-run{background:transparent;border-color:var(--cyan);color:var(--cyan);text-shadow:0 0 8px var(--cyan);box-shadow:0 0 12px rgba(0,245,255,.15),inset 0 0 12px rgba(0,245,255,.05)}
  .btn-run:hover:not(:disabled){background:rgba(0,245,255,.08);box-shadow:0 0 20px rgba(0,245,255,.35)}
  .btn-run:disabled{border-color:var(--dim);color:var(--dim);cursor:not-allowed;box-shadow:none}
  .btn-teal{background:transparent;border-color:var(--teal);color:var(--teal)}
  .btn-teal:hover{background:rgba(0,200,212,.08);box-shadow:0 0 12px rgba(0,200,212,.3)}
  .btn-green{background:transparent;border-color:var(--green);color:var(--green)}
  .btn-green:hover{background:rgba(0,255,65,.08);box-shadow:0 0 12px rgba(0,255,65,.3)}
  .btn-ghost{background:transparent;border-color:var(--border);color:var(--text-dim)}
  .btn-ghost:hover{border-color:var(--magenta);color:var(--magenta)}
  .btn-spin{position:relative}
  .btn-spin.loading::after{content:'';position:absolute;inset:0;background:rgba(0,0,0,.4)}

  /* Status messages */
  .msg{margin-top:10px;font-size:.78rem;padding:8px 12px;border-radius:2px;display:none}
  .msg.ok {background:rgba(0,255,65,.07);border:1px solid rgba(0,255,65,.3);color:var(--green);display:block}
  .msg.err{background:rgba(255,34,68,.07);border:1px solid rgba(255,34,68,.3);color:var(--red);display:block}
  .msg.info{background:rgba(0,245,255,.05);border:1px solid rgba(0,245,255,.2);color:var(--cyan);display:block}

  /* Table */
  .table-wrap{overflow-x:auto;margin-top:4px;border:1px solid var(--border);border-radius:2px}
  table{width:100%;border-collapse:collapse;font-size:.78rem}
  thead th{padding:9px 14px;font-size:.62rem;letter-spacing:.18em;color:var(--cyan);background:var(--bg3);border-bottom:1px solid var(--border);text-align:left;white-space:nowrap}
  tbody td{padding:8px 14px;border-bottom:1px solid rgba(26,26,58,.5);color:var(--text);vertical-align:top}
  tbody tr:last-child td{border-bottom:none}
  tbody tr:hover td{background:rgba(0,245,255,.02)}
  .td-n{color:var(--text-dim);font-size:.7rem;width:36px}
  .td-id{color:var(--magenta2);font-size:.72rem;white-space:nowrap}
  .td-dest{color:var(--cyan);font-family:inherit}
  .td-dur,.td-delay{color:var(--amber)}
  .td-tags{color:var(--magenta2);font-size:.72rem}
  .td-trunc{font-size:.72rem;color:var(--text-dim);max-width:200px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}

  /* Filename row */
  .fname-row{display:flex;align-items:center;gap:10px;margin-bottom:12px}
  .fname-row label{white-space:nowrap;font-size:.68rem;letter-spacing:.15em;color:var(--text-dim);margin:0}
  .fname-row input{flex:1;background:var(--bg1);border:1px solid var(--border);color:var(--cyan);font-family:inherit;font-size:.82rem;padding:7px 10px;border-radius:2px;outline:none;margin:0;transition:border-color .2s}
  .fname-row input:focus{border-color:var(--cyan)}

  /* Script preview */
  .script-pre{background:var(--bg1);border:1px solid var(--border);border-radius:4px;padding:16px;font-size:.77rem;line-height:1.65;white-space:pre;overflow:auto;max-height:420px;font-family:'Courier New',monospace;margin:0 0 12px;scrollbar-width:thin;scrollbar-color:var(--border) transparent}
  .script-pre::-webkit-scrollbar{width:5px;height:5px}
  .script-pre::-webkit-scrollbar-thumb{background:var(--border);border-radius:3px}

  /* Robot syntax highlight */
  .r-sec{color:var(--cyan);text-shadow:0 0 6px rgba(0,245,255,.45)}
  .r-name{color:var(--magenta);text-shadow:0 0 5px rgba(255,0,255,.35)}
  .r-set{color:var(--amber)}
  .r-var{color:var(--green)}
  .r-cmt{color:#334466}
  .r-kw{color:#c8c8e8}

  /* Footer */
  .footer{display:flex;align-items:center;gap:12px;margin-top:20px;font-size:.72rem;color:var(--text-dim);letter-spacing:.08em;flex-wrap:wrap}
  .footer-right{margin-left:auto}
  .hidden{display:none!important}
</style>
</head>
<body>
<div class="wrap">

  <header>
    <div style="display:flex;align-items:center;gap:14px">
      <a href="/" class="back-link">&#8592; CONTROL PANEL</a>
      <div class="logo">SCRIPT<span>//</span>GENERATOR</div>
    </div>
    <div class="header-meta">
      <div class="pill">EXCEL &#8594; ROBOT</div>
      <div class="pill online">&#11044; ONLINE</div>
    </div>
  </header>

  <!-- ── STEP 1: UPLOAD ── -->
  <div class="card">
    <div class="card-title">// 01 &middot; UPLOAD EXCEL FILE</div>

    <div class="drop-zone" id="dropZone">
      <div class="drop-icon">&#11014;</div>
      <div class="drop-text">DRAG &amp; DROP YOUR .XLSX FILE HERE</div>
      <div class="drop-sub">Supports .xlsx format &middot; first row must be column headers</div>
    </div>
    <input type="file" id="fileInput" accept=".xlsx,.xls" style="display:none">

    <div style="margin-top:14px">
      <div style="font-size:.68rem;letter-spacing:.15em;color:var(--text-dim);margin-bottom:8px">TEMPLATE TYPE</div>
      <div class="type-row">
        <label class="type-opt">
          <input type="radio" name="ttype" value="call" checked>
          <span>&#128222; CALL TEST &nbsp;—&nbsp; ADB mobile dial &amp; end</span>
        </label>
        <label class="type-opt">
          <input type="radio" name="ttype" value="tms">
          <span>&#128196; TMS TEST CASES &nbsp;—&nbsp; TMSID / Headline / Steps format</span>
        </label>
      </div>
    </div>

    <div class="btn-row">
      <button class="btn-teal" onclick="downloadTemplate()">&#8595; DOWNLOAD TEMPLATE</button>
      <button class="btn-run"  onclick="document.getElementById('fileInput').click()">&#128193; BROWSE FILE</button>
    </div>

    <div class="file-info hidden" id="fileInfo">
      <span style="color:var(--green)">&#10003;</span>
      <span class="file-name" id="fileName"></span>
      <span class="file-size" id="fileSize"></span>
    </div>
    <div class="msg" id="uploadMsg"></div>
  </div>

  <!-- ── STEP 2: PREVIEW TABLE ── -->
  <div class="card hidden" id="previewCard">
    <div class="card-title">
      // 02 &middot; PARSED TEST CASES
      <span id="testCount" style="color:var(--magenta);margin-left:8px"></span>
    </div>
    <div class="table-wrap">
      <table>
        <thead id="previewHead"></thead>
        <tbody id="previewBody"></tbody>
      </table>
    </div>
  </div>

  <!-- ── STEP 3: GENERATED SCRIPT ── -->
  <div class="card hidden" id="scriptCard">
    <div class="card-title">// 03 &middot; GENERATED ROBOT SCRIPT</div>
    <div class="fname-row">
      <label>OUTPUT FILE</label>
      <input type="text" id="outputFilename" value="generated_calls.robot">
    </div>
    <div class="script-pre" id="scriptPre"></div>
    <div class="btn-row">
      <button class="btn-run"   onclick="saveScript()">&#128190; SAVE SCRIPT</button>
      <button class="btn-green" onclick="runScript()">&#9654; RUN NOW</button>
      <button class="btn-ghost" onclick="copyScript()">&#128203; COPY</button>
    </div>
    <div class="msg" id="scriptMsg"></div>
  </div>

  <div class="footer">
    <span id="footerTxt">AWAITING INPUT</span>
    <div class="footer-right" id="footer-time"></div>
  </div>
</div>

<script>
'use strict';

// ── Clock ──────────────────────────────────────────────────────────
function tick(){
  var d=new Date();
  document.getElementById('footer-time').textContent=
    d.toLocaleTimeString('en-GB',{hour12:false})+'  UTC'+
    (d.getTimezoneOffset()<=0?'+':'-')+
    String(Math.abs(d.getTimezoneOffset()/60)).padStart(2,'0');
}
setInterval(tick,1000);tick();

// ── State ──────────────────────────────────────────────────────────
var parsedScript='';
var detectedFmt='call';

// ── Drag & Drop ────────────────────────────────────────────────────
var dz=document.getElementById('dropZone');
dz.addEventListener('click',function(){document.getElementById('fileInput').click();});
dz.addEventListener('dragover',function(e){e.preventDefault();dz.classList.add('over');});
dz.addEventListener('dragleave',function(){dz.classList.remove('over');});
dz.addEventListener('drop',function(e){
  e.preventDefault();dz.classList.remove('over');
  var f=e.dataTransfer.files[0];if(f)handleFile(f);
});
document.getElementById('fileInput').addEventListener('change',function(e){
  if(e.target.files[0])handleFile(e.target.files[0]);
});

// ── File handling ──────────────────────────────────────────────────
function handleFile(file){
  if(!file.name.match(/\.xlsx?$/i)){
    showMsg('uploadMsg','Only .xlsx / .xls files are supported.','err');return;
  }
  document.getElementById('fileName').textContent=file.name;
  document.getElementById('fileSize').textContent=fmtSize(file.size);
  document.getElementById('fileInfo').classList.remove('hidden');
  document.getElementById('footerTxt').textContent='PARSING '+file.name.toUpperCase()+' …';
  showMsg('uploadMsg','Reading file…','info');
  var reader=new FileReader();
  reader.onload=function(ev){
    var b64=ev.target.result.split(',')[1];
    parseExcel(b64,file.name);
  };
  reader.readAsDataURL(file);
}

function fmtSize(b){
  if(b<1024)return b+' B';
  if(b<1048576)return (b/1024).toFixed(1)+' KB';
  return (b/1048576).toFixed(1)+' MB';
}

// ── Parse Excel ────────────────────────────────────────────────────
async function parseExcel(b64,fname){
  var fmt=document.querySelector('input[name="ttype"]:checked').value;
  try{
    var res=await fetch('/parse-excel',{
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify({data:b64,filename:fname,fmt:fmt})
    });
    var data=await res.json();
    if(data.error){showMsg('uploadMsg','✗ '+data.error,'err');return;}
    showMsg('uploadMsg','✓ Parsed '+data.count+' test case(s) — format: '+data.format.toUpperCase(),'ok');
    detectedFmt=data.format;
    // Suggest filename
    var base=fname.replace(/\.xlsx?$/i,'');
    document.getElementById('outputFilename').value=base+'_generated.robot';
    renderTable(data.tests,data.format);
    renderScript(data.script);
    document.getElementById('footerTxt').textContent='READY — '+data.count+' TEST CASE(S) LOADED';
  }catch(e){
    showMsg('uploadMsg','✗ Network error: '+e,'err');
  }
}

// ── Table rendering ────────────────────────────────────────────────
function esc(s){return String(s).replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;')}
function trunc(s,n){s=String(s||'');return s.length>n?s.slice(0,n)+'…':s;}

function renderTable(tests,fmt){
  document.getElementById('testCount').textContent='[ '+tests.length+' ENTRIES ]';
  var isTms=fmt==='tms';
  var thead;
  if(isTms){
    thead='<tr><th>#</th><th>TMSID</th><th>HEADLINE / TEST NAME</th><th>DESCRIPTION</th><th>STEPS (PREVIEW)</th></tr>';
  } else {
    thead='<tr><th>#</th><th>TEST CASE NAME</th><th>DESTINATION</th><th>DURATION</th><th>DELAY</th><th>TAGS</th><th>DOCUMENTATION</th></tr>';
  }
  document.getElementById('previewHead').innerHTML=thead;
  var rows=tests.map(function(t,i){
    if(isTms){
      return '<tr>'+
        '<td class="td-n">'+(i+1)+'</td>'+
        '<td class="td-id">'+esc(t.tmsid||'')+'</td>'+
        '<td>'+esc(t.name)+'</td>'+
        '<td class="td-trunc" title="'+esc(t.description||'')+'">'+esc(trunc(t.description,80))+'</td>'+
        '<td class="td-trunc" title="'+esc(t.steps||'')+'">'+esc(trunc(t.steps,80))+'</td>'+
        '</tr>';
    } else {
      return '<tr>'+
        '<td class="td-n">'+(i+1)+'</td>'+
        '<td>'+esc(t.name)+'</td>'+
        '<td class="td-dest">'+esc(t.destination||'')+'</td>'+
        '<td class="td-dur">'+esc(t.duration||'')+'</td>'+
        '<td class="td-delay">'+esc(t.delay||'')+'</td>'+
        '<td class="td-tags">'+esc(t.tags||'')+'</td>'+
        '<td class="td-trunc">'+esc(trunc(t.doc,60))+'</td>'+
        '</tr>';
    }
  }).join('');
  document.getElementById('previewBody').innerHTML=rows;
  document.getElementById('previewCard').classList.remove('hidden');
}

// ── Script syntax highlight ────────────────────────────────────────
function highlight(text){
  return text.split('\n').map(function(line){
    var e=line.replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;');
    if(/^\*{3}/.test(line))           return '<span class="r-sec">'+e+'</span>';
    if(/^#/.test(line.trim()))         return '<span class="r-cmt">'+e+'</span>';
    if(/^\s{4}\[/.test(line))          return '<span class="r-set">'+e+'</span>';
    if(/^\S/.test(line)&&line.trim()&&!/^#/.test(line)) return '<span class="r-name">'+e+'</span>';
    e=e.replace(/(\$\{[^}]+\}|@\{[^}]+\})/g,'<span class="r-var">$1</span>');
    return '<span class="r-kw">'+e+'</span>';
  }).join('\n');
}

function renderScript(script){
  parsedScript=script;
  document.getElementById('scriptPre').innerHTML=highlight(script);
  document.getElementById('scriptCard').classList.remove('hidden');
}

// ── Save ───────────────────────────────────────────────────────────
async function saveScript(){
  var filename=document.getElementById('outputFilename').value.trim()||'generated.robot';
  try{
    var res=await fetch('/save-script',{
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify({filename:filename,content:parsedScript})
    });
    var data=await res.json();
    if(data.ok){
      showMsg('scriptMsg','✓ Saved to '+data.path,'ok');
      document.getElementById('footerTxt').textContent='SAVED: '+filename.toUpperCase();
    } else {
      showMsg('scriptMsg','✗ '+(data.error||'Save failed'),'err');
    }
  }catch(e){showMsg('scriptMsg','✗ '+e,'err');}
}

// ── Copy ───────────────────────────────────────────────────────────
async function copyScript(){
  try{
    await navigator.clipboard.writeText(parsedScript);
    showMsg('scriptMsg','✓ Copied to clipboard','ok');
  }catch{showMsg('scriptMsg','✗ Clipboard access denied','err');}
}

// ── Run Now ────────────────────────────────────────────────────────
async function runScript(){
  var filename=document.getElementById('outputFilename').value.trim()||'generated.robot';
  try{
    await fetch('/save-script',{
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify({filename:filename,content:parsedScript})
    });
    window.location.href='/?suite='+encodeURIComponent(filename);
  }catch(e){showMsg('scriptMsg','✗ '+e,'err');}
}

// ── Download template ──────────────────────────────────────────────
function downloadTemplate(){
  var fmt=document.querySelector('input[name="ttype"]:checked').value;
  window.location.href='/download-template?fmt='+fmt;
}

// ── Utilities ──────────────────────────────────────────────────────
function showMsg(id,msg,cls){
  var el=document.getElementById(id);
  el.textContent=msg;
  el.className='msg '+cls;
}
</script>
</body>
</html>
"""


# ═══════════════════════════════════════════════════════════════════════════════
#  BACKEND — ADB & ROBOT
# ═══════════════════════════════════════════════════════════════════════════════

def get_adb_devices():
    adb = ADB_PATH if os.path.exists(ADB_PATH) else "adb"
    version, devices = "", []
    try:
        r = subprocess.run([adb, "version"], capture_output=True, text=True, timeout=5)
        for line in r.stdout.splitlines():
            if "Android Debug Bridge" in line:
                version = line.strip(); break
    except Exception as e:
        version = f"adb error: {e}"
    try:
        r = subprocess.run([adb, "devices"], capture_output=True, text=True, timeout=5)
        for line in r.stdout.splitlines()[1:]:
            parts = line.split()
            if len(parts) >= 2:
                serial = parts[0]
                devices.append({"serial": serial, "status": parts[1],
                                 "name": DEVICE_NAMES.get(serial, "")})
    except Exception:
        pass
    return version, devices


def run_robot(suite, phone):
    global _proc
    with state_lock:
        run_state["running"] = True
        run_state["output"]  = []
        run_state["returncode"] = None
        run_state["done"] = False
    env = os.environ.copy()
    env["PHONE_NUMBER"] = phone
    try:
        _proc = subprocess.Popen(
            ["python", "-m", "robot", suite],
            cwd=ROBOT_DIR, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
            text=True, bufsize=1, env=env,
        )
        for line in _proc.stdout:
            with state_lock:
                run_state["output"].append(line.rstrip())
        _proc.wait()
        with state_lock:
            run_state["returncode"] = _proc.returncode
    except Exception as e:
        with state_lock:
            run_state["output"].append(f"ERROR: {e}")
            run_state["returncode"] = 1
    finally:
        with state_lock:
            run_state["running"] = False
            run_state["done"] = True
        _proc = None


# ═══════════════════════════════════════════════════════════════════════════════
#  BACKEND — SCRIPT GENERATOR
# ═══════════════════════════════════════════════════════════════════════════════

def list_robot_scripts():
    try:
        return sorted(f for f in os.listdir(ROBOT_DIR) if f.endswith('.robot'))
    except Exception:
        return []


def _rf_escape(text):
    """Escape RF variable syntax so Log statements show literal text."""
    return str(text).replace('${', r'\${').replace('@{', r'\@{').replace('%{', r'\%{')


def _split_numbered_steps(text):
    """Split text on numbered step patterns like '1.' '2.' etc."""
    parts = re.split(r'(?=\n?\d+\.)', str(text).strip())
    result = []
    for p in parts:
        p = re.sub(r'\s+', ' ', p.strip().replace('\r', ''))
        if p:
            result.append(p[:250])
    return result or [str(text).strip()[:250]]


def generate_call_script(tests):
    today = datetime.date.today().isoformat()
    count = len(tests)
    lines = [
        f'# Generated by openitt Test Automation Control Panel',
        f'# Date      : {today}',
        f'# Format    : Call Test (ADB)',
        f'# Test cases: {count}',
        '',
        '*** Settings ***',
        f'Documentation    Generated call test suite — {count} test case(s).',
        'Library          OperatingSystem',
        'Library          Process',
        'Library          String',
        '',
        '*** Variables ***',
        '${ADB_TIMEOUT}    30',
        '',
        '*** Keywords ***',
        'Verify ADB Is Available',
        '    ${result}=    Run Process    adb    --version    stdout=${OUTPUT_DIR}/adb_version.txt',
        '    Should Be Equal As Integers    ${result.rc}    0',
        '    ${output}=    Get File    ${OUTPUT_DIR}/adb_version.txt',
        '    Should Contain    ${output}    Android Debug Bridge',
        '',
        'Verify Device Is Connected',
        '    ${result}=    Run Process    adb    devices    stdout=${OUTPUT_DIR}/devices.txt',
        '    Should Be Equal As Integers    ${result.rc}    0',
        '    ${output}=    Get File    ${OUTPUT_DIR}/devices.txt',
        '    Should Match Regexp    ${output}    (?m)^\\S+\\s+device$',
        '    Log    Device connected successfully    level=INFO',
        '',
        'Dial Phone Number',
        '    [Arguments]    ${number}',
        '    Log    Dialing ${number} ...    level=INFO',
        '    ${result}=    Run Process    adb    shell    am    start    -a    android.intent.action.CALL    -d    tel:${number}',
        '    ...    timeout=${ADB_TIMEOUT}s',
        '    Should Be Equal As Integers    ${result.rc}    0',
        '    Log    Dial command executed successfully    level=INFO',
        '',
        'Wait And End Call',
        '    [Arguments]    ${duration}=5s',
        '    Sleep    ${duration}',
        '    ${result}=    Run Process    adb    shell    input    keyevent    6',
        '    ...    timeout=${ADB_TIMEOUT}s',
        '    Should Be Equal As Integers    ${result.rc}    0',
        '    Log    Call ended    level=INFO',
        '',
        '*** Test Cases ***',
        'Setup And Verify Environment',
        '    [Documentation]    Verify ADB is installed and a device is connected.',
        '    Verify ADB Is Available',
        '    Verify Device Is Connected',
        '',
    ]
    for i, t in enumerate(tests):
        lines.append(t['name'])
        if t.get('doc'):
            lines.append(f"    [Documentation]    {t['doc']}")
        if t.get('tags'):
            lines.append(f"    [Tags]    {t['tags']}")
        lines.append(f"    Dial Phone Number    {t['destination']}")
        lines.append(f"    Wait And End Call    {t['duration']}")
        if i < count - 1 and t.get('delay', '0s') not in ('0s', '0'):
            lines.append(f"    Sleep    {t['delay']}")
        lines.append('')
    return '\n'.join(lines)


def generate_tms_script(tests):
    today = datetime.date.today().isoformat()
    count = len(tests)
    lines = [
        f'# Generated by openitt Test Automation Control Panel',
        f'# Date      : {today}',
        f'# Format    : TMS Test Cases',
        f'# Test cases: {count}',
        '',
        '*** Settings ***',
        f'Documentation    TMS test suite — {count} test case(s) generated {today}.',
        '...    Review and implement step keywords before executing.',
        'Library    OperatingSystem',
        'Library    String',
        'Library    Collections',
        '',
        '*** Keywords ***',
        'Log Step',
        '    [Arguments]    ${step}',
        '    Log    ${step}    level=INFO    console=yes',
        '',
        '*** Test Cases ***',
    ]
    for t in tests:
        lines.append(t['name'])
        # Documentation block
        doc_lines = []
        if t.get('tmsid'):
            doc_lines.append(f"TMS-ID: {t['tmsid']}")
        if t.get('description'):
            for chunk in textwrap.wrap(t['description'], 90):
                doc_lines.append(chunk)
        if t.get('expected'):
            doc_lines.append('Expected:')
            for chunk in textwrap.wrap(t['expected'], 90):
                doc_lines.append(chunk)
        if doc_lines:
            lines.append(f'    [Documentation]    {doc_lines[0]}')
            for dl in doc_lines[1:]:
                lines.append(f'    ...    {dl}')
        # Tags
        tag_parts = ['tms']
        if t.get('tmsid'):
            tag_parts.append(t['tmsid'])
        lines.append(f'    [Tags]    ' + '    '.join(tag_parts))
        # Pre-conditions
        if t.get('preconditions'):
            lines.append('    # --- Pre-Conditions ---')
            for step in _split_numbered_steps(t['preconditions']):
                lines.append(f'    Log Step    {_rf_escape(step)}')
        # Test steps
        if t.get('steps'):
            lines.append('    # --- Test Steps ---')
            for step in _split_numbered_steps(t['steps']):
                lines.append(f'    Log Step    {_rf_escape(step)}')
        if t.get('comments'):
            lines.append(f'    # Note: {_rf_escape(t["comments"][:120])}')
        lines.append('')
    return '\n'.join(lines)


def parse_excel_and_generate(excel_bytes, fmt):
    if not HAS_OPENPYXL:
        return {'error': 'openpyxl is not installed — run: pip install openpyxl'}
    try:
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb.active
    except Exception as e:
        return {'error': f'Cannot read Excel file: {e}'}

    # Read headers from row 1 to auto-detect format
    headers = [str(c.value or '').strip().lower() for c in ws[1]]
    if fmt == 'auto':
        fmt = 'tms' if any(h in ('tmsid', 'headline', 'tms id') for h in headers) else 'call'

    tests = []
    if fmt == 'call':
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(v for v in row if v is not None):
                continue
            name = str(row[0] or '').strip() if len(row) > 0 else ''
            dest = str(row[1] or '').strip() if len(row) > 1 else ''
            if not name or not dest:
                continue
            dur_raw = row[2] if len(row) > 2 else None
            del_raw = row[3] if len(row) > 3 else None
            tags    = str(row[4] or 'mobile,call').strip() if len(row) > 4 else 'mobile,call'
            doc     = str(row[5] or '').strip()            if len(row) > 5 else ''
            dur = str(dur_raw).strip() if dur_raw is not None else '5'
            if not dur.endswith('s'): dur += 's'
            dly = str(del_raw).strip() if del_raw is not None else '3'
            if not dly.endswith('s'): dly += 's'
            tests.append({'name': name, 'destination': dest, 'duration': dur,
                          'delay': dly, 'tags': tags, 'doc': doc})
        script = generate_call_script(tests)

    else:  # tms
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(v for v in row if v is not None):
                continue
            tmsid   = str(row[0] or '').strip() if len(row) > 0 else ''
            name    = str(row[1] or '').strip() if len(row) > 1 else ''
            if not name:
                continue
            desc    = str(row[2] or '').strip() if len(row) > 2 else ''
            precond = str(row[3] or '').strip() if len(row) > 3 else ''
            steps   = str(row[4] or '').strip() if len(row) > 4 else ''
            expected= str(row[5] or '').strip() if len(row) > 5 else ''
            comments= str(row[6] or '').strip() if len(row) > 6 else ''
            tests.append({'tmsid': tmsid, 'name': name, 'description': desc,
                          'preconditions': precond, 'steps': steps,
                          'expected': expected, 'comments': comments})
        script = generate_tms_script(tests)

    return {'tests': tests, 'script': script, 'count': len(tests), 'format': fmt}


def save_robot_script(filename, content):
    filename = os.path.basename(filename)
    if not filename.endswith('.robot'):
        filename += '.robot'
    path = os.path.join(ROBOT_DIR, filename)
    with open(path, 'w', encoding='utf-8') as f:
        f.write(content)
    return path


def create_excel_template(fmt):
    if not HAS_OPENPYXL:
        return None
    wb  = openpyxl.Workbook()
    ws  = wb.active
    hdr_fill = PatternFill(fill_type='solid', fgColor='0D3D4E')
    hdr_font = Font(bold=True, color='00F5FF')
    row_fill = PatternFill(fill_type='solid', fgColor='0D0D1A')
    row_font = Font(color='C8C8E8')

    if fmt == 'tms':
        ws.title = 'TMS Test Cases'
        headers  = ['TMSID', 'Headline', 'Description', 'Pre-Conditions', 'Test Steps', 'ExpectedResults', 'Comments']
        widths   = {'A':18,'B':35,'C':40,'D':40,'E':50,'F':50,'G':30}
        examples = [
            ['TMSII01276126', 'Voice call OnNet charge < 60s',
             'Verify CCS charges subscriber for voice onnet call.',
             '1. Matrixx RT-CCS is working fine.\n2. Subscription has sufficient balance.',
             '1. Query wallet from REST interface.\n2. Initiate voice call.\n3. Disconnect after 50s.\n4. Verify EDR.',
             '1. Wallet queried successfully.\n2. Call connected.\n3. 0.03€ charged.\n4. EDR verified.',
             'AldiTalk subscription required'],
            ['TMSII01276150', 'Data usage charge from main balance',
             'Verify CCS charges subscriber for data usage from main balance.',
             '1. Matrixx RT-CCS working.\n2. Sufficient balance.',
             '1. Query wallet.\n2. Initiate data session.\n3. Verify EDR.\n4. Verify charge.',
             '1. Wallet queried.\n2. Data session started.\n3. EDR verified.\n4. Charge verified.',
             ''],
        ]
    else:
        ws.title = 'Call Tests'
        headers  = ['Test Case Name', 'Destination Number', 'Duration (s)', 'Delay After (s)', 'Tags', 'Documentation']
        widths   = {'A':28,'B':22,'C':15,'D':15,'E':22,'F':40}
        examples = [
            ['Call to Alice',     '+49151234567',   5, 3, 'mobile,call',           'Test call to Alice'],
            ['Call to Bob',       '+49152345678',  10, 5, 'mobile,call',           'Test call to Bob'],
            ['Emergency Line',    '+4916012345678', 5, 3, 'mobile,emergency',      'Emergency line check'],
            ['Support Hotline',   '+49800123456',   8, 3, 'mobile,call,support',   'Support queue test'],
        ]

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = hdr_fill; c.font = hdr_font
        c.alignment = Alignment(horizontal='center', vertical='center')

    for ri, row in enumerate(examples, 2):
        for ci, val in enumerate(row, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.fill = row_fill; c.font = row_font

    col_letters = [chr(65+i) for i in range(len(headers))]
    for letter in col_letters:
        if letter in widths:
            ws.column_dimensions[letter].width = widths[letter]

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════════
#  HTTP HANDLER
# ═══════════════════════════════════════════════════════════════════════════════

_proc = None


def serve_file(handler, path, content_type):
    if os.path.exists(path):
        with open(path, "rb") as f:
            data = f.read()
        handler.send_response(200)
        handler.send_header("Content-Type", content_type)
        handler.send_header("Content-Length", str(len(data)))
        handler.end_headers()
        handler.wfile.write(data)
    else:
        handler.send_response(404)
        handler.end_headers()
        handler.wfile.write(b"File not found")


def json_response(handler, obj, code=200):
    payload = json.dumps(obj).encode()
    handler.send_response(code)
    handler.send_header("Content-Type", "application/json")
    handler.send_header("Content-Length", str(len(payload)))
    handler.end_headers()
    handler.wfile.write(payload)


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        pass  # suppress noise

    def do_GET(self):
        parsed = urlparse(self.path)
        path   = parsed.path
        qs     = parse_qs(parsed.query)

        if path == "/":
            data = HTML.encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)

        elif path == "/generator":
            data = GENERATOR_HTML.encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)

        elif path == "/status":
            from_idx = int(qs.get("from", ["0"])[0])
            with state_lock:
                lines = run_state["output"][from_idx:]
                done  = run_state["done"]
                rc    = run_state["returncode"]
            json_response(self, {"lines": lines, "done": done, "returncode": rc})

        elif path == "/devices":
            version, devices = get_adb_devices()
            json_response(self, {"version": version, "devices": devices})

        elif path == "/list-scripts":
            json_response(self, {"scripts": list_robot_scripts()})

        elif path == "/download-template":
            fmt  = qs.get("fmt", ["call"])[0]
            data = create_excel_template(fmt)
            if data is None:
                self.send_response(503)
                self.end_headers()
                self.wfile.write(b"openpyxl not installed - run: pip install openpyxl")
                return
            fname = f"openitt_{fmt}_template.xlsx"
            self.send_response(200)
            self.send_header("Content-Type",
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", f'attachment; filename="{fname}"')
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)

        elif path == "/report":
            serve_file(self, os.path.join(ROBOT_DIR, "report.html"), "text/html")

        elif path == "/log":
            serve_file(self, os.path.join(ROBOT_DIR, "log.html"), "text/html")

        else:
            self.send_response(404)
            self.end_headers()

    def do_POST(self):
        length = int(self.headers.get("Content-Length", 0))
        body   = json.loads(self.rfile.read(length) or b"{}") if length else {}

        if self.path == "/run":
            suite = body.get("suite", "call_test.robot")
            phone = body.get("phone", "555")
            with state_lock:
                already = run_state["running"]
            if not already:
                threading.Thread(target=run_robot, args=(suite, phone), daemon=True).start()
            json_response(self, {"status": "started"})

        elif self.path == "/abort":
            global _proc
            if _proc:
                try:
                    _proc.terminate()
                except Exception:
                    pass
            json_response(self, {"status": "aborted"})

        elif self.path == "/parse-excel":
            excel_bytes = base64.b64decode(body.get("data", ""))
            fmt         = body.get("fmt", "call")
            result      = parse_excel_and_generate(excel_bytes, fmt)
            json_response(self, result)

        elif self.path == "/save-script":
            filename = body.get("filename", "generated.robot")
            content  = body.get("content", "")
            try:
                path = save_robot_script(filename, content)
                json_response(self, {"ok": True, "path": path})
            except Exception as e:
                json_response(self, {"ok": False, "error": str(e)})

        else:
            self.send_response(404)
            self.end_headers()


# ═══════════════════════════════════════════════════════════════════════════════
#  ENTRY POINT
# ═══════════════════════════════════════════════════════════════════════════════

def main():
    server = HTTPServer(("localhost", PORT), Handler)
    url    = f"http://localhost:{PORT}"
    print(f"\n  TESTAUTOMATION CONTROL PANEL")
    print(f"  ─────────────────────────────")
    print(f"  URL       : {url}")
    print(f"  Generator : {url}/generator")
    openpyxl_status = "available" if HAS_OPENPYXL else "MISSING — run: pip install openpyxl"
    print(f"  openpyxl  : {openpyxl_status}")
    print(f"  Press Ctrl+C to stop\n")
    webbrowser.open(url)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n  Server stopped.")


if __name__ == "__main__":
    main()
